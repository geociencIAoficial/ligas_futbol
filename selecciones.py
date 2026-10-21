#!/usr/bin/env python3
"""
Base de datos de convocatorias y participaciones de selecciones nacionales.

Única fuente: Wikipedia, vía su API oficial (MediaWiki API). Gratis, sin
registro ni llave. Contenido bajo licencia CC BY-SA 4.0.
  - Listas de 26 del Mundial 2026 ("2026 FIFA World Cup squads").
  - Alineaciones, cambios y tarjetas de los 104 partidos (páginas de cada grupo
    y de la fase final), de donde se calculan los minutos jugados.
  - "Current squad" de cada selección, para fotografiar las convocatorias de
    cada fecha FIFA posterior al Mundial.

No hace scraping de sitios que lo prohíben (Transfermarkt, Sofascore, FotMob...).

Uso:
    python selecciones.py init                 # pide tu correo, una sola vez
    python selecciones.py mundial              # listas de 26
    python selecciones.py partidos             # alineaciones y minutos del Mundial
    python selecciones.py ligas                # liga de cada club (repetir tras cada snapshot)
    python selecciones.py tablas               # tabla de posiciones de cada liga
    python selecciones.py panel                # panel.html: clasificación mundial de clubes
    python selecciones.py snapshot --ventana 2026-10
    python selecciones.py reporte                  # Excel con una hoja por selección
    python selecciones.py reporte --ventana 2026-10  # además, la hoja Recambio
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sqlite3
import time
import unicodedata
from datetime import date, datetime, timezone
from pathlib import Path

import requests
from bs4 import BeautifulSoup, NavigableString, Tag

# ---------------------------------------------------------------------------
# Configuración
# ---------------------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "selecciones.db"
CACHE_DIR = BASE_DIR / "cache"
CONFIG_PATH = BASE_DIR / "config.json"

WIKI_API = "https://en.wikipedia.org/w/api.php"
WIKI_PAUSA_S = 1.0          # una petición por segundo, de sobra para su política

WC_SQUADS = "2026_FIFA_World_Cup_squads"
WC_PARTIDOS = [f"2026_FIFA_World_Cup_Group_{g}" for g in "ABCDEFGHIJKL"] + [
    "2026_FIFA_World_Cup_knockout_stage"]
WC_ETIQUETA = "Mundial 2026"

POSICIONES = {"GK", "DF", "MF", "FW"}
# Posiciones que Wikipedia usa en las alineaciones de partido
POS_PARTIDO = re.compile(r"^(GK|RB|CB|LB|RWB|LWB|SW|DF|DM|CM|LM|RM|AM|MF|RW|LW|"
                         r"RF|LF|CF|SS|ST|FW|RCB|LCB|RCM|LCM|RAM|LAM)$")

SCHEMA = """
CREATE TABLE IF NOT EXISTS selecciones (
    seleccion_id  INTEGER PRIMARY KEY,
    nombre        TEXT UNIQUE NOT NULL,       -- nombre en Wikipedia (inglés)
    wiki_pagina   TEXT,                       -- p.ej. Mexico_national_football_team
    categoria     TEXT                        -- '1 Élite' ... '4 Resto' (ver CAT_SELECCIONES)
);
CREATE TABLE IF NOT EXISTS jugadores (
    jugador_id    INTEGER PRIMARY KEY,
    wiki_titulo   TEXT UNIQUE,                -- identificador estable
    nombre        TEXT NOT NULL,
    fecha_nac     TEXT
);
CREATE TABLE IF NOT EXISTS clubes (
    club_id       INTEGER PRIMARY KEY,
    nombre        TEXT NOT NULL,
    pais_club     TEXT,                       -- asociación a la que pertenece el club
    wiki_titulo   TEXT,
    liga          TEXT,                       -- del infobox del club ('' si no se encontró)
    categoria     TEXT,                       -- '1 Top 10 mundial' ... '4 Resto' (ver CAT_CLUBES)
    UNIQUE(nombre, pais_club)
);
-- Una lista = una convocatoria publicada (Mundial, fecha FIFA de octubre, etc.)
CREATE TABLE IF NOT EXISTS listas (
    lista_id      INTEGER PRIMARY KEY,
    seleccion_id  INTEGER NOT NULL REFERENCES selecciones,
    ventana       TEXT NOT NULL,              -- 'mundial-2026', '2026-10', ...
    fuente        TEXT NOT NULL,
    nota          TEXT,
    tecnico       TEXT,
    obtenida_en   TEXT NOT NULL,
    UNIQUE(seleccion_id, ventana)
);
CREATE TABLE IF NOT EXISTS convocados (
    lista_id      INTEGER NOT NULL REFERENCES listas,
    jugador_id    INTEGER NOT NULL REFERENCES jugadores,
    dorsal        INTEGER,
    posicion      TEXT,
    partidos_sel  INTEGER,                    -- caps al momento de la lista
    goles_sel     INTEGER,
    club_id       INTEGER REFERENCES clubes,  -- club al momento de la lista
    capitan       INTEGER DEFAULT 0,
    PRIMARY KEY (lista_id, jugador_id)
);
CREATE TABLE IF NOT EXISTS partidos (
    partido_id    INTEGER PRIMARY KEY,
    competicion   TEXT,
    ronda         TEXT,                       -- 'Group A', 'Round of 16', ...
    fecha         TEXT,
    local_id      INTEGER REFERENCES selecciones,
    visita_id     INTEGER REFERENCES selecciones,
    marcador      TEXT,
    prorroga      INTEGER DEFAULT 0,
    fuente        TEXT,
    UNIQUE(fecha, local_id, visita_id)
);
CREATE TABLE IF NOT EXISTS participaciones (
    partido_id    INTEGER NOT NULL REFERENCES partidos,
    seleccion_id  INTEGER NOT NULL REFERENCES selecciones,
    jugador_id    INTEGER NOT NULL REFERENCES jugadores,
    dorsal        INTEGER,
    posicion      TEXT,
    titular       INTEGER,                    -- 1 titular, 0 banca
    min_entrada   INTEGER,                    -- NULL si no entró
    min_salida    INTEGER,
    minutos       INTEGER,
    amarillas     INTEGER DEFAULT 0,
    expulsado     INTEGER DEFAULT 0,
    PRIMARY KEY (partido_id, jugador_id)
);
"""


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------
def conectar() -> sqlite3.Connection:
    con = sqlite3.connect(DB_PATH)
    # Bases creadas con la versión anterior (API-Football) tienen otras columnas
    # en partidos/participaciones; esas tablas se recrean (las listas se conservan).
    cols = {r[1] for r in con.execute("PRAGMA table_info(partidos)")}
    if cols and "local_id" not in cols:
        con.executescript("DROP TABLE IF EXISTS participaciones; DROP TABLE IF EXISTS partidos;")
        print("Actualicé las tablas de partidos al formato nuevo.")
    elif cols and "ronda" not in cols:
        con.execute("ALTER TABLE partidos ADD COLUMN ronda TEXT")
        print("Agregué la columna 'ronda'. Vuelve a correr 'partidos' para llenarla.")
    for tabla, col in (("clubes", "liga"), ("clubes", "categoria"), ("selecciones", "categoria")):
        existentes = {r[1] for r in con.execute(f"PRAGMA table_info({tabla})")}
        if existentes and col not in existentes:
            con.execute(f"ALTER TABLE {tabla} ADD COLUMN {col} TEXT")
    con.execute("PRAGMA foreign_keys = ON")
    con.executescript(SCHEMA)   # crea las tablas si aún no existen
    con.row_factory = sqlite3.Row
    return con


def norm(texto: str | None) -> str:
    """Minúsculas, sin acentos ni signos: para comparar nombres."""
    if not texto:
        return ""
    t = unicodedata.normalize("NFKD", texto)
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = re.sub(r"[^a-zA-Z0-9 ]+", " ", t).lower()
    return re.sub(r"\s+", " ", t).strip()


def a_entero(texto: str | None) -> int | None:
    m = re.search(r"\d+", texto or "")
    return int(m.group()) if m else None


def cache_get(clave: str) -> dict | None:
    f = CACHE_DIR / f"{hashlib.sha1(clave.encode()).hexdigest()}.json"
    return json.loads(f.read_text("utf-8")) if f.exists() else None


def cache_put(clave: str, datos: dict) -> None:
    CACHE_DIR.mkdir(exist_ok=True)
    f = CACHE_DIR / f"{hashlib.sha1(clave.encode()).hexdigest()}.json"
    f.write_text(json.dumps(datos, ensure_ascii=False), "utf-8")


def ahora() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


# ---------------------------------------------------------------------------
# Configuración del usuario (se pide una sola vez y se guarda en config.json)
# ---------------------------------------------------------------------------
def config_contacto() -> str:
    if os.environ.get("WIKI_CONTACTO"):           # GitHub Actions: viene de un secreto
        return os.environ["WIKI_CONTACTO"]
    datos = json.loads(CONFIG_PATH.read_text("utf-8")) if CONFIG_PATH.exists() else {}
    if datos.get("contacto"):
        return datos["contacto"]
    valor = ""
    while "@" not in valor:
        valor = input("Tu correo electrónico (Wikipedia pide que quien consulta su API deje\n"
                      "un contacto; no es una cuenta ni una contraseña): ").strip()
        if "@" not in valor:
            print("  Eso no parece un correo, inténtalo de nuevo.")
    datos["contacto"] = valor
    CONFIG_PATH.write_text(json.dumps(datos, indent=2), "utf-8")
    print(f"  Guardado en {CONFIG_PATH.name}.")
    return valor


def user_agent() -> str:
    # Política de Wikipedia: https://meta.wikimedia.org/wiki/User-Agent_policy
    return f"SeleccionesDB/0.2 (proyecto personal de análisis; {config_contacto()})"


# ---------------------------------------------------------------------------
# Wikipedia (API oficial)
# ---------------------------------------------------------------------------
_wiki = requests.Session()


def wiki_html(pagina: str, refrescar: bool = False) -> str:
    """HTML renderizado de una página, vía action=parse. Con caché por día."""
    clave = f"wiki:{pagina}:{date.today().isoformat()}"
    if not refrescar and (c := cache_get(clave)):
        return c["html"]
    _wiki.headers["User-Agent"] = user_agent()
    params = {
        "action": "parse", "page": pagina, "prop": "text",
        "format": "json", "formatversion": 2, "redirects": 1,
        "maxlag": 5,  # se retira solo si los servidores están cargados
    }
    for intento in range(4):
        r = _wiki.get(WIKI_API, params=params, timeout=30)
        datos = r.json()
        if datos.get("error", {}).get("code") == "maxlag":
            time.sleep(5 * (intento + 1))
            continue
        if "error" in datos:
            raise RuntimeError(f"Wikipedia ({pagina}): {datos['error'].get('info')}")
        break
    time.sleep(WIKI_PAUSA_S)
    html = datos["parse"]["text"]
    cache_put(clave, {"html": html})
    return html


def _texto(el: Tag) -> str:
    return re.sub(r"\s+", " ", el.get_text(" ", strip=True)).strip()


def _raiz(html: str) -> Tag:
    sopa = BeautifulSoup(html, "lxml")
    return sopa.find(class_="mw-parser-output") or sopa


def _enlace_jugador(celda: Tag) -> Tag | None:
    """Primer enlace a un artículo que no sea bandera, archivo ni nota."""
    for a in celda.find_all("a", title=True):
        if a.find_parent(class_="flagicon") or a.find("img"):
            continue
        href = a.get("href", "")
        if href.startswith("#") or "File:" in href or not a.get_text(strip=True):
            continue
        return a
    return None


def titulo(a: Tag) -> str:
    """Título del artículo enlazado (sin la marca de enlace rojo)."""
    return a["title"].replace(" (page does not exist)", "").strip()


# ---------------------------------------------------------------------------
# Parseo: tablas de plantilla (listas de convocados)
# ---------------------------------------------------------------------------
def _es_tabla_plantilla(tabla: Tag) -> bool:
    encabezado = " ".join(_texto(th) for th in tabla.find_all("th")[:10]).lower()
    return "pos" in encabezado and "player" in encabezado and "club" in encabezado


def _celda_club(td: Tag) -> tuple[str, str | None, str | None]:
    """Devuelve (nombre_club, pais_del_club, titulo_wiki_club)."""
    pais = None
    flag = td.find(class_="flagicon")
    if flag:
        a = flag.find("a", title=True)
        img = flag.find("img")
        pais = (a["title"] if a else None) or (img.get("alt") if img else None)
        flag.extract()
    enlaces = [a for a in td.find_all("a", title=True) if not a["href"].startswith("#")]
    if enlaces:
        return _texto(enlaces[-1]), pais, enlaces[-1]["title"]
    return _texto(td), pais, None


def _fecha_nac(td: Tag) -> str | None:
    bday = td.find(class_="bday")
    if bday:
        return bday.get_text(strip=True)
    texto = re.sub(r"\(.*?\)", "", _texto(td)).strip()
    for fmt in ("%d %B %Y", "%B %d, %Y"):
        try:
            return datetime.strptime(texto, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def parsear_tabla_plantilla(tabla: Tag) -> list[dict]:
    """Columnas: No. | Pos. | Player | Date of birth (age) | Caps | Goals | Club"""
    filas = tabla.find_all("tr")
    encabezados = [norm(_texto(c)) for c in filas[0].find_all(["th", "td"])]

    def col(*claves):
        for i, h in enumerate(encabezados):
            if any(h.startswith(k) for k in claves):
                return i
        return None

    i_no, i_pos, i_jug = col("no"), col("pos"), col("player")
    i_nac, i_caps, i_gol, i_club = col("date of birth"), col("caps"), col("goals"), col("club")

    jugadores = []
    for tr in filas[1:]:
        celdas = tr.find_all(["td", "th"])
        if i_jug is None or len(celdas) <= max(i for i in (i_jug, i_club) if i is not None):
            continue
        c_jug = celdas[i_jug]
        enlace = _enlace_jugador(c_jug)
        nombre = _texto(enlace) if enlace else re.sub(r"\(.*?\)", "", _texto(c_jug)).strip()
        if not nombre:
            continue
        pos_txt = _texto(celdas[i_pos]) if i_pos is not None else ""
        pos = next((p for p in POSICIONES if re.search(rf"\b{p}\b", pos_txt)), None)
        club, pais_club, club_wiki = (_celda_club(celdas[i_club]) if i_club is not None
                                      else (None, None, None))
        jugadores.append({
            "nombre": nombre,
            "wiki_titulo": titulo(enlace) if enlace else None,
            "dorsal": a_entero(_texto(celdas[i_no])) if i_no is not None else None,
            "posicion": pos,
            "fecha_nac": _fecha_nac(celdas[i_nac]) if i_nac is not None else None,
            "partidos_sel": a_entero(_texto(celdas[i_caps])) if i_caps is not None else None,
            "goles_sel": a_entero(_texto(celdas[i_gol])) if i_gol is not None else None,
            "club": club, "pais_club": pais_club, "club_wiki": club_wiki,
            "capitan": int("captain" in _texto(c_jug).lower()),
        })
    return jugadores


def _encabezado(el: Tag) -> tuple[int, str, str] | None:
    """Si el elemento es un título de sección, devuelve (nivel, texto, id)."""
    h = None
    if el.name in ("h2", "h3", "h4"):
        h = el
    elif el.name == "div" and "mw-heading" in (el.get("class") or []):
        h = el.find(["h2", "h3", "h4"])
    if h is None:
        return None
    titulo = h.get_text(" ", strip=True).replace("[edit]", "").strip()
    return int(h.name[1]), titulo, h.get("id", "")


def plantillas_mundial(html: str) -> list[dict]:
    """Grupo (h2) > selección (h3) > párrafos y tabla de 26."""
    equipos, actual = [], None
    for el in _raiz(html).find_all(["div", "h2", "h3", "p", "table"], recursive=False):
        enc = _encabezado(el)
        if enc:
            nivel, titulo, _ = enc
            if nivel == 3:
                actual = {"seleccion": titulo, "tecnico": None,
                          "wiki_pagina": None, "nota": "", "jugadores": []}
                equipos.append(actual)
            elif nivel == 2:
                actual = None
            continue
        if actual is None:
            continue
        if el.name == "p":
            texto = _texto(el)
            if texto.lower().startswith(("coach", "head coach", "manager")):
                a = el.find("a", title=True)
                actual["tecnico"] = _texto(a) if a else texto.split(":", 1)[-1].strip()
            else:
                actual["nota"] += texto + " "
            if not actual["wiki_pagina"]:
                for a in el.find_all("a", href=True):
                    if re.search(r"national_(soccer_|football_)?team", a["href"]):
                        actual["wiki_pagina"] = a["href"].split("/wiki/")[-1]
                        break
        elif el.name == "table" and _es_tabla_plantilla(el) and not actual["jugadores"]:
            actual["jugadores"] = parsear_tabla_plantilla(el)
    return [e for e in equipos if e["jugadores"]]


def plantilla_actual(html: str) -> dict | None:
    """Extrae la sección 'Current squad' de la página de una selección."""
    dentro, nivel_sec, nota = False, None, ""
    for el in _raiz(html).find_all(["div", "h2", "h3", "h4", "p", "table"], recursive=False):
        enc = _encabezado(el)
        if enc:
            nivel, titulo, ident = enc
            if dentro and nivel <= nivel_sec:
                return None
            if "current squad" in titulo.lower() or ident.lower().startswith("current_squad"):
                dentro, nivel_sec = True, nivel
            continue
        if not dentro:
            continue
        if el.name == "p":
            nota += _texto(el) + " "
        elif el.name == "table" and _es_tabla_plantilla(el):
            return {"nota": nota.strip(), "jugadores": parsear_tabla_plantilla(el)}
    return None


# ---------------------------------------------------------------------------
# Parseo: partidos (Football box + tablas de alineación)
# ---------------------------------------------------------------------------
def _minuto(texto: str) -> int | None:
    """'71′' -> 71 ; '90+3′' -> 90 ; '105+1' -> 105."""
    m = re.search(r"(\d{1,3})(?:\s*\+\s*\d+)?", texto)
    return int(m.group(1)) if m else None


def _tipo_evento(img: Tag) -> str | None:
    padre = img.find_parent(title=True)
    s = " ".join([img.get("alt", ""), img.get("src", ""), padre.get("title", "") if padre else ""])
    s = s.lower().replace("_", " ")
    if "sub off" in s or "substituted off" in s or "downward" in s:
        return "sale"
    if "sub on" in s or "substituted on" in s or "upward" in s:
        return "entra"
    if "yellow-red" in s or "second yellow" in s or "red card" in s or "sent off" in s:
        return "roja"
    if "yellow card" in s or "booked" in s:
        return "amarilla"
    return None


def _eventos_fila(tr: Tag, despues_de: Tag) -> list[tuple[str, int | None]]:
    """Eventos (tipo, minuto) de una fila, leyendo las imágenes tras el nombre."""
    eventos = []
    imgs = [i for i in tr.find_all("img") if despues_de not in i.parents]
    for img in imgs:
        tipo = _tipo_evento(img)
        if not tipo:
            continue
        minuto = None
        for el in img.next_elements:
            if isinstance(el, Tag) and el.name == "img":
                break
            if isinstance(el, NavigableString) and (m := _minuto(str(el))) is not None:
                minuto = m
                break
            if isinstance(el, Tag) and el.name == "tr":
                break
        eventos.append((tipo, minuto))
    return eventos


def _es_tabla_alineacion(tabla: Tag) -> bool:
    texto = tabla.get_text(" ", strip=True)
    if not re.search(r"Substitut", texto):
        return False
    # que no contenga a su vez otra alineación (nos quedamos con la tabla interna)
    for t in tabla.find_all("table"):
        if re.search(r"Substitut", t.get_text(" ", strip=True)):
            return False
    filas_jugador = 0
    for tr in tabla.find_all("tr"):
        celdas = tr.find_all(["td", "th"])
        if celdas and POS_PARTIDO.match(_texto(celdas[0])):
            filas_jugador += 1
    return filas_jugador >= 11


def parsear_alineacion(tabla: Tag, duracion: int) -> tuple[list[dict], str | None]:
    """Devuelve (jugadores, técnico). Calcula minutos jugados de cada uno."""
    jugadores, banca, tecnico, en_tecnico = [], False, None, False
    for tr in tabla.find_all("tr"):
        texto = _texto(tr)
        if re.match(r"^Substitut", texto):
            banca = True
            continue
        if re.match(r"^(Manager|Head coach|Coach)", texto):
            en_tecnico = True
            continue
        if en_tecnico:
            a = tr.find("a", title=True)
            tecnico = _texto(a) if a else texto
            en_tecnico = False
            continue
        celdas = tr.find_all(["td", "th"])
        if len(celdas) < 3 or not POS_PARTIDO.match(_texto(celdas[0])):
            continue
        dorsal = next((int(_texto(c)) for c in celdas[1:3] if _texto(c).isdigit()), None)
        enlace, celda_nombre = None, None
        for c in celdas[1:]:
            enlace = _enlace_jugador(c)
            if enlace:
                celda_nombre = c
                break
        if not enlace:
            continue
        ev = _eventos_fila(tr, celda_nombre)
        sale = min((m for t, m in ev if t in ("sale", "roja") and m is not None), default=None)
        entra = next((m for t, m in ev if t == "entra" and m is not None), None)
        if not banca:
            entrada, salida = 0, (sale if sale is not None else duracion)
        elif entra is not None:
            entrada, salida = entra, (sale if sale is not None and sale >= entra else duracion)
        else:
            entrada, salida = None, None
        jugadores.append({
            "nombre": _texto(enlace), "wiki_titulo": titulo(enlace), "fecha_nac": None,
            "dorsal": dorsal, "posicion": _texto(celdas[0]), "titular": int(not banca),
            "min_entrada": entrada, "min_salida": salida,
            "minutos": (max(salida - entrada, 1) if entrada is not None else 0),
            "amarillas": sum(1 for t, _ in ev if t == "amarilla"),
            "expulsado": int(any(t == "roja" for t, _ in ev)),
        })
    return jugadores, tecnico


def _equipo_box(th: Tag | None) -> tuple[str | None, str | None]:
    """(nombre visible, título del artículo de la selección) de un lado del marcador."""
    if th is None:
        return None, None
    a = next((a for a in th.find_all("a", title=True) if not a.find_parent(class_="flagicon")), None)
    if a:
        return _texto(a), a["title"]
    return _texto(th), None


def partidos_de_pagina(html: str) -> list[dict]:
    """Recorre la página en orden: cada Football box abre un partido y las dos
    tablas de alineación que le siguen son la local y la visitante."""
    partidos, actual, ronda = [], None, None
    raiz = _raiz(html)
    for el in raiz.find_all(lambda t: t.name in ("table", "h2") or
                            (t.name == "div" and "footballbox" in (t.get("class") or []))):
        if el.name == "h2":
            ronda = el.get_text(" ", strip=True).replace("[edit]", "").strip()
            continue
        if el.name == "div":
            texto = _texto(el)
            fecha = None
            f = el.find(class_=re.compile(r"dtstart|bday"))
            if f and re.match(r"\d{4}-\d{2}-\d{2}", f.get_text(strip=True)):
                fecha = f.get_text(strip=True)[:10]
            local, local_t = _equipo_box(el.find(class_="fhome"))
            visita, visita_t = _equipo_box(el.find(class_="faway"))
            score = el.find(class_="fscore")
            actual = {"fecha": fecha, "ronda": ronda, "local": local, "local_titulo": local_t,
                      "visita": visita, "visita_titulo": visita_t,
                      "marcador": _texto(score) if score else None,
                      "prorroga": int("a.e.t" in texto or "extra time" in texto.lower()),
                      "tablas": []}
            partidos.append(actual)
        elif actual is not None and len(actual["tablas"]) < 2 and _es_tabla_alineacion(el):
            if el.find_parent(class_="footballbox"):
                continue
            actual["tablas"].append(el)
    salida = []
    for p in partidos:
        dur = 120 if p["prorroga"] else 90
        lados = [parsear_alineacion(t, dur) for t in p.pop("tablas")]
        if len(lados) == 2:
            p["alineacion_local"], p["tecnico_local"] = lados[0]
            p["alineacion_visita"], p["tecnico_visita"] = lados[1]
        else:
            p["alineacion_local"] = p["alineacion_visita"] = None  # Wikipedia no la publica
        salida.append(p)
    return salida


# ---------------------------------------------------------------------------
# Escritura en la base
# ---------------------------------------------------------------------------
def upsert_seleccion(con, nombre, wiki_pagina=None) -> int:
    con.execute("INSERT OR IGNORE INTO selecciones(nombre) VALUES (?)", (nombre,))
    if wiki_pagina:
        con.execute("UPDATE selecciones SET wiki_pagina=? WHERE nombre=? AND wiki_pagina IS NULL",
                    (wiki_pagina, nombre))
    return con.execute("SELECT seleccion_id FROM selecciones WHERE nombre=?", (nombre,)).fetchone()[0]


def buscar_seleccion(con, nombre, titulo) -> int | None:
    """Ubica una selección por el artículo enlazado o, si no, por su nombre."""
    if titulo:
        r = con.execute("SELECT seleccion_id FROM selecciones WHERE wiki_pagina=?",
                        (titulo.replace(" ", "_"),)).fetchone()
        if r:
            return r[0]
    if nombre:
        for r in con.execute("SELECT seleccion_id, nombre FROM selecciones"):
            if norm(r["nombre"]) == norm(nombre):
                return r[0]
    return None


def upsert_jugador(con, j) -> int:
    if j.get("wiki_titulo"):
        con.execute("INSERT OR IGNORE INTO jugadores(wiki_titulo, nombre, fecha_nac) VALUES (?,?,?)",
                    (j["wiki_titulo"], j["nombre"], j.get("fecha_nac")))
        if j.get("fecha_nac"):
            con.execute("UPDATE jugadores SET fecha_nac=COALESCE(fecha_nac, ?) WHERE wiki_titulo=?",
                        (j["fecha_nac"], j["wiki_titulo"]))
        return con.execute("SELECT jugador_id FROM jugadores WHERE wiki_titulo=?",
                           (j["wiki_titulo"],)).fetchone()[0]
    fila = con.execute("SELECT jugador_id FROM jugadores WHERE wiki_titulo IS NULL AND nombre=? "
                       "AND COALESCE(fecha_nac,'')=COALESCE(?, '')",
                       (j["nombre"], j.get("fecha_nac"))).fetchone()
    if fila:
        return fila[0]
    return con.execute("INSERT INTO jugadores(nombre, fecha_nac) VALUES (?,?)",
                       (j["nombre"], j.get("fecha_nac"))).lastrowid


def upsert_club(con, j) -> int | None:
    if not j["club"]:
        return None
    con.execute("INSERT OR IGNORE INTO clubes(nombre, pais_club, wiki_titulo) VALUES (?,?,?)",
                (j["club"], j["pais_club"], j["club_wiki"]))
    return con.execute("SELECT club_id FROM clubes WHERE nombre=? AND pais_club IS ?",
                       (j["club"], j["pais_club"])).fetchone()[0]


def guardar_lista(con, seleccion_id, ventana, fuente, nota, tecnico, jugadores) -> int:
    con.execute("DELETE FROM convocados WHERE lista_id IN "
                "(SELECT lista_id FROM listas WHERE seleccion_id=? AND ventana=?)",
                (seleccion_id, ventana))
    con.execute("DELETE FROM listas WHERE seleccion_id=? AND ventana=?", (seleccion_id, ventana))
    lista_id = con.execute(
        "INSERT INTO listas(seleccion_id, ventana, fuente, nota, tecnico, obtenida_en) "
        "VALUES (?,?,?,?,?,?)",
        (seleccion_id, ventana, fuente, (nota or "").strip()[:1000], tecnico, ahora())).lastrowid
    for j in jugadores:
        con.execute("INSERT OR REPLACE INTO convocados VALUES (?,?,?,?,?,?,?,?)",
                    (lista_id, upsert_jugador(con, j), j["dorsal"], j["posicion"],
                     j["partidos_sel"], j["goles_sel"], upsert_club(con, j), j["capitan"]))
    return lista_id


def guardar_partido(con, p: dict, competicion: str, fuente: str) -> int | None:
    local = buscar_seleccion(con, p["local"], p["local_titulo"])
    visita = buscar_seleccion(con, p["visita"], p["visita_titulo"])
    if not local or not visita:
        print(f"  No reconozco a {p['local']} o {p['visita']}; partido omitido.")
        return None
    con.execute("INSERT INTO partidos(competicion, ronda, fecha, local_id, visita_id, marcador,"
                " prorroga, fuente) VALUES (?,?,?,?,?,?,?,?) "
                "ON CONFLICT(fecha, local_id, visita_id) DO UPDATE SET competicion=excluded.competicion,"
                " ronda=excluded.ronda, marcador=excluded.marcador, prorroga=excluded.prorroga,"
                " fuente=excluded.fuente",
                (competicion, p.get("ronda"), p["fecha"], local, visita, p["marcador"],
                 p["prorroga"], fuente))
    pid = con.execute("SELECT partido_id FROM partidos WHERE fecha IS ? AND local_id=? AND visita_id=?",
                      (p["fecha"], local, visita)).fetchone()[0]
    if p["alineacion_local"] is None:      # partido registrado, pero sin alineaciones
        return pid
    con.execute("DELETE FROM participaciones WHERE partido_id=?", (pid,))
    for sel, lista in ((local, p["alineacion_local"]), (visita, p["alineacion_visita"])):
        for j in lista:
            con.execute("INSERT OR REPLACE INTO participaciones VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                        (pid, sel, upsert_jugador(con, j), j["dorsal"], j["posicion"], j["titular"],
                         j["min_entrada"], j["min_salida"], j["minutos"],
                         j["amarillas"], j["expulsado"]))
    return pid


# ---------------------------------------------------------------------------
# Comandos
# ---------------------------------------------------------------------------
def cmd_init(_args):
    conectar().close()
    print(f"Base lista en {DB_PATH}")
    config_contacto()
    print("Siguiente paso: python selecciones.py mundial")


def cmd_mundial(args):
    equipos = plantillas_mundial(wiki_html(WC_SQUADS, args.refrescar))
    fuente = f"https://en.wikipedia.org/wiki/{WC_SQUADS}"
    with conectar() as con:
        for e in equipos:
            sid = upsert_seleccion(con, e["seleccion"], e["wiki_pagina"])
            guardar_lista(con, sid, "mundial-2026", fuente, e["nota"], e["tecnico"], e["jugadores"])
    total = sum(len(e["jugadores"]) for e in equipos)
    print(f"{len(equipos)} selecciones y {total} jugadores cargados.")
    if len(equipos) != 48:
        print("Aviso: se esperaban 48 selecciones; revisa si cambió el formato de la página.")


def cmd_partidos(args):
    paginas = args.paginas or WC_PARTIDOS
    competicion = args.competicion or (WC_ETIQUETA if not args.paginas else "Otra")
    with conectar() as con:
        if not con.execute("SELECT 1 FROM selecciones LIMIT 1").fetchone():
            print("Primero corre: python selecciones.py mundial")
            return
        total, sin_total = 0, []
        for pagina in paginas:
            try:
                partidos = partidos_de_pagina(wiki_html(pagina, args.refrescar))
            except Exception as ex:
                print(f"  {pagina}: error {ex}")
                continue
            if m := re.search(r"Group_([A-Z])$", pagina):
                for p in partidos:          # en las páginas de grupo la ronda es el grupo
                    p["ronda"] = f"Group {m.group(1)}"
            guardados, sin_alin = 0, []
            for p in partidos:
                if guardar_partido(con, p, competicion, f"https://en.wikipedia.org/wiki/{pagina}"):
                    if p["alineacion_local"] is None:
                        sin_alin.append(p)
                    else:
                        guardados += 1
            total += guardados
            sin_total.extend(sin_alin)
            extra = f" (+{len(sin_alin)} sin alineación en Wikipedia)" if sin_alin else ""
            print(f"  {pagina.replace('_', ' ')}: {guardados} partidos{extra}")
            con_alin = [p for p in partidos if p["alineacion_local"] is not None]
            if args.muestra and con_alin:
                p = con_alin[0]
                print(f"\n  Muestra: {p['local']} {p['marcador']} {p['visita']} ({p['fecha']})")
                for j in p["alineacion_local"]:
                    print(f"    {'T' if j['titular'] else 'S'} #{j['dorsal']:<3}{j['posicion']:<4}"
                          f"{j['nombre']:<28}{j['minutos']:>4}'")
                print()
        # control de calidad: cada partido debe tener 11 titulares por lado
        raros = con.execute("""SELECT p.fecha, s.nombre, COUNT(*) n FROM participaciones pa
                               JOIN partidos p USING(partido_id) JOIN selecciones s
                               ON s.seleccion_id = pa.seleccion_id
                               WHERE pa.titular=1 GROUP BY pa.partido_id, pa.seleccion_id
                               HAVING n <> 11""").fetchall()
    print(f"{total} partidos con alineación guardados" +
          (" (el Mundial tuvo 104)." if not args.paginas else "."))
    if sin_total:
        print(f"{len(sin_total)} partidos quedaron registrados sin alineación, porque Wikipedia "
              "no la publica:")
        for p in sin_total:
            print(f"  {p['fecha']}  {p['local']} {p['marcador']} {p['visita']}")
    for r in raros:
        print(f"  Revisar: {r['nombre']} el {r['fecha']} tiene {r['n']} titulares")


def cmd_snapshot(args):
    """Guarda la 'Current squad' de cada selección como la lista de una ventana FIFA."""
    with conectar() as con:
        filtro = [s.lower() for s in args.equipos] if args.equipos else None
        sels = [s for s in con.execute("SELECT * FROM selecciones WHERE wiki_pagina IS NOT NULL")
                if not filtro or s["nombre"].lower() in filtro]
        sin_tabla = []
        for s in sels:
            try:
                res = plantilla_actual(wiki_html(s["wiki_pagina"], args.refrescar))
            except Exception as ex:
                print(f"  {s['nombre']}: error {ex}")
                continue
            if not res or not res["jugadores"]:
                sin_tabla.append(s["nombre"])
                continue
            guardar_lista(con, s["seleccion_id"], args.ventana,
                          f"https://en.wikipedia.org/wiki/{s['wiki_pagina']}",
                          res["nota"], None, res["jugadores"])
            print(f"  {s['nombre']}: {len(res['jugadores'])} convocados")
        if sin_tabla:
            print("Sin sección 'Current squad' legible:", ", ".join(sin_tabla))
    print("Revisa en la nota de cada lista (tabla 'listas') para qué partidos fue la "
          "convocatoria; Wikipedia la actualiza a mano y a veces va atrasada.")


def _edad(fecha_nac: str | None, ref: date) -> float | None:
    if not fecha_nac:
        return None
    return (ref - date.fromisoformat(fecha_nac)).days / 365.25


NOMBRES_ES = {
    "Algeria": "Argelia", "Argentina": "Argentina", "Australia": "Australia", "Austria": "Austria",
    "Belgium": "Bélgica", "Bosnia and Herzegovina": "Bosnia y Herzegovina", "Brazil": "Brasil",
    "Canada": "Canadá", "Cape Verde": "Cabo Verde", "Colombia": "Colombia", "Croatia": "Croacia",
    "Curaçao": "Curazao", "Czech Republic": "Chequia", "DR Congo": "RD Congo", "Ecuador": "Ecuador",
    "Egypt": "Egipto", "England": "Inglaterra", "France": "Francia", "Germany": "Alemania",
    "Ghana": "Ghana", "Haiti": "Haití", "Iran": "Irán", "Iraq": "Irak", "Ivory Coast": "Costa de Marfil",
    "Japan": "Japón", "Jordan": "Jordania", "Mexico": "México", "Morocco": "Marruecos",
    "Netherlands": "Países Bajos", "New Zealand": "Nueva Zelanda", "Norway": "Noruega",
    "Panama": "Panamá", "Paraguay": "Paraguay", "Portugal": "Portugal", "Qatar": "Catar",
    "Saudi Arabia": "Arabia Saudita", "Scotland": "Escocia", "Senegal": "Senegal",
    "South Africa": "Sudáfrica", "South Korea": "Corea del Sur", "Spain": "España",
    "Sweden": "Suecia", "Switzerland": "Suiza", "Tunisia": "Túnez", "Turkey": "Turquía",
    "United States": "Estados Unidos", "Uruguay": "Uruguay", "Uzbekistan": "Uzbekistán",
    "Greece": "Grecia", "Russia": "Rusia", "Cyprus": "Chipre", "Romania": "Rumania",
    "Bulgaria": "Bulgaria", "Hungary": "Hungría", "Slovakia": "Eslovaquia", "Slovenia": "Eslovenia",
    "Finland": "Finlandia", "Ireland": "Irlanda", "Republic of Ireland": "Irlanda",
    "Northern Ireland": "Irlanda del Norte", "Israel": "Israel", "Kazakhstan": "Kazajistán",
    "Azerbaijan": "Azerbaiyán", "Armenia": "Armenia", "Georgia": "Georgia",
    "United Arab Emirates": "Emiratos Árabes Unidos", "Kuwait": "Kuwait", "Bahrain": "Baréin",
    "Oman": "Omán", "Lebanon": "Líbano", "China": "China", "Malaysia": "Malasia",
    "Indonesia": "Indonesia", "Thailand": "Tailandia", "India": "India", "Libya": "Libia",
    "Italy": "Italia", "Denmark": "Dinamarca", "Poland": "Polonia", "Wales": "Gales",
    "Serbia": "Serbia", "Ukraine": "Ucrania", "Nigeria": "Nigeria", "Cameroon": "Camerún",
    "Peru": "Perú", "Chile": "Chile", "Venezuela": "Venezuela", "Bolivia": "Bolivia",
    "Costa Rica": "Costa Rica", "Honduras": "Honduras", "Jamaica": "Jamaica",
    "El Salvador": "El Salvador", "Guatemala": "Guatemala",
}
RONDAS_ES = {
    "round of 32": "Dieciseisavos", "round of 16": "Octavos", "quarter-finals": "Cuartos",
    "quarterfinals": "Cuartos", "semi-finals": "Semifinal", "semifinals": "Semifinal",
    "third place play-off": "Tercer lugar", "third place match": "Tercer lugar",
    "match for third place": "Tercer lugar", "final": "Final",
}


def nombre_es(nombre: str) -> str:
    return NOMBRES_ES.get(nombre, nombre)


def ronda_es(ronda: str | None) -> str:
    if not ronda:
        return ""
    if m := re.match(r"Group ([A-Z])$", ronda):
        return f"Grupo {m.group(1)}"
    return RONDAS_ES.get(ronda.strip().lower(), ronda)


def _ventana_de_partido(competicion: str, fecha: str | None) -> str | None:
    if competicion == WC_ETIQUETA:
        return "mundial-2026"
    return fecha[:7] if fecha else None


def datos_recambio(con, ventana: str) -> list[dict]:
    ref_wc, ref_v = date(2026, 6, 11), date.today()
    filas = []
    for s in con.execute("SELECT * FROM selecciones ORDER BY nombre").fetchall():
        def lista(v):
            return {r["jugador_id"]: r for r in con.execute(
                """SELECT j.jugador_id, j.nombre, j.fecha_nac FROM convocados c
                   JOIN listas l USING(lista_id) JOIN jugadores j USING(jugador_id)
                   WHERE l.seleccion_id=? AND l.ventana=?""", (s["seleccion_id"], v))}
        a, b = lista("mundial-2026"), lista(ventana)
        if not a or not b:
            continue
        siguen, salen, entran = a.keys() & b.keys(), a.keys() - b.keys(), b.keys() - a.keys()
        edades_a = [e for e in (_edad(r["fecha_nac"], ref_wc) for r in a.values()) if e]
        edades_b = [e for e in (_edad(r["fecha_nac"], ref_v) for r in b.values()) if e]
        mins = [m for m in con.execute(
            """SELECT j.fecha_nac, SUM(pa.minutos) m FROM participaciones pa
               JOIN partidos p USING(partido_id) JOIN jugadores j USING(jugador_id)
               WHERE p.competicion=? AND pa.seleccion_id=? GROUP BY j.jugador_id""",
            (WC_ETIQUETA, s["seleccion_id"])) if m["fecha_nac"] and m["m"]]
        tot = sum(m["m"] for m in mins)
        pond = sum(_edad(m["fecha_nac"], ref_wc) * m["m"] for m in mins) / tot if tot else None
        min_salen = sum(r[0] or 0 for r in con.execute(
            f"""SELECT SUM(pa.minutos) FROM participaciones pa JOIN partidos p USING(partido_id)
                WHERE p.competicion=? AND pa.seleccion_id=?
                AND pa.jugador_id IN ({','.join('?' * len(salen)) or 'NULL'})""",
            (WC_ETIQUETA, s["seleccion_id"], *salen)))
        filas.append({
            "seleccion": nombre_es(s["nombre"]), "categoria": categoria_seleccion(s["nombre"]),
            "mundial": len(a), "ventana": len(b),
            "siguen": len(siguen), "salen": len(salen), "entran": len(entran),
            "edad_a": round(sum(edades_a) / len(edades_a), 1) if edades_a else None,
            "edad_pond": round(pond, 1) if pond else None,
            "edad_b": round(sum(edades_b) / len(edades_b), 1) if edades_b else None,
            "min_salen": min_salen,
            "salen_n": "; ".join(sorted((a[i]["nombre"] for i in salen), key=norm)),
            "entran_n": "; ".join(sorted((b[i]["nombre"] for i in entran), key=norm)),
        })
    filas.sort(key=lambda f: f["siguen"] / f["mundial"])
    return filas


# ---------------------------------------------------------------------------
# Categorías de selecciones y clubes
# ---------------------------------------------------------------------------
# Son criterios propios, pensados para agrupar y no para ordenar con precisión.
# Puedes editarlos: el script recalcula las categorías cada vez que genera el reporte.
#
# Selecciones (nombres en inglés, como en Wikipedia):
#   1 Élite: campeones del mundo recientes o semifinalistas habituales, con plantillas
#     llenas de jugadores de clubes de élite.
#   2 Segunda línea: llegan con regularidad a eliminatorias del Mundial o ganan su
#     confederación, pero rara vez pelean el título.
#   3 Buen equipo: clasifican seguido a Mundiales o finales continentales.
#   4 Resto: debutantes, clasificaciones esporádicas o selecciones en formación.
CAT_SELECCIONES = {
    "1 Élite": ["Argentina", "France", "Spain", "England", "Brazil", "Germany", "Portugal",
                "Netherlands"],
    "2 Segunda línea": ["Belgium", "Croatia", "Uruguay", "Italy", "Colombia", "Morocco",
                        "Switzerland", "Denmark", "Japan", "Mexico", "United States", "Senegal"],
    "3 Buen equipo": ["Ecuador", "Austria", "Turkey", "Norway", "Sweden", "South Korea", "Iran",
                      "Australia", "Paraguay", "Ivory Coast", "Egypt", "Algeria", "Canada",
                      "Scotland", "Czech Republic", "Tunisia", "Nigeria", "Poland", "Serbia",
                      "Ukraine", "Wales", "Peru", "Chile", "Ghana", "Cameroon"],
}
CAT_SEL_RESTO = "4 Resto"

# Clubes:
#   1 Top 10 mundial: los que combinan títulos de Champions recientes, coeficiente UEFA
#     de los últimos cinco años, presupuesto e historia.
#   2 Top segunda línea: candidatos habituales a cuartos de Champions o gigantes
#     históricos de Sudamérica.
#   3 Importante: cualquier club de primera división de las cinco grandes ligas, más
#     los clubes dominantes de otras ligas fuertes (lista de abajo).
#   4 Resto.
CAT_CLUBES = {
    "1 Top 10 mundial": [
        "Real Madrid", "Barcelona", "Bayern Munich", "Manchester City", "Liverpool",
        "Paris Saint-Germain", "Arsenal", "Inter Milan", "Chelsea", "Manchester United"],
    "2 Top segunda línea": [
        "Atlético Madrid", "Juventus", "AC Milan", "Borussia Dortmund", "Bayer Leverkusen",
        "Napoli", "Tottenham Hotspur", "Newcastle United", "Aston Villa", "Atalanta",
        "Benfica", "Porto", "Sporting CP", "RB Leipzig", "Roma", "Ajax", "PSV Eindhoven",
        "Marseille", "Monaco", "Flamengo", "Palmeiras", "River Plate", "Boca Juniors"],
    "3 Importante": [
        "Feyenoord", "Galatasaray", "Fenerbahçe", "Beşiktaş", "Celtic", "Rangers",
        "Club Brugge", "Anderlecht", "Olympiacos", "Panathinaikos", "PAOK",
        "Shakhtar Donetsk", "Dinamo Zagreb", "Red Bull Salzburg", "Copenhagen", "Young Boys",
        "Slavia Prague", "Sparta Prague", "Red Star Belgrade", "Braga", "AZ", "Twente",
        "Al-Hilal", "Al-Nassr", "Al-Ittihad", "Al-Ahli", "Al Ahly", "Zamalek", "Pyramids",
        "Mamelodi Sundowns", "Espérance de Tunis", "Corinthians", "São Paulo", "Santos",
        "Botafogo", "Fluminense", "Grêmio", "Internacional", "Atlético Mineiro", "Cruzeiro",
        "Vasco da Gama", "Racing Club", "Independiente", "Estudiantes", "Peñarol", "Nacional",
        "América", "Guadalajara", "Cruz Azul", "Monterrey", "Tigres UANL", "Pumas UNAM",
        "Toluca", "Pachuca", "Inter Miami CF", "LA Galaxy", "Los Angeles FC", "Al Ain",
        "Ulsan HD", "Jeonbuk Hyundai Motors", "Kashima Antlers", "Urawa Red Diamonds"],
}
CAT_CLUB_RESTO = "4 Resto"
# Primeras divisiones de las cinco grandes ligas (título del artículo en Wikipedia)
LIGAS_TOP5 = {"Premier League", "La Liga", "Serie A", "Bundesliga", "Ligue 1"}

_PALABRAS_VACIAS = {"fc", "cf", "sc", "afc", "sk", "fk", "cd", "ca", "club", "de", "futbol",
                    "football", "calcio", "1907", "1909", "1899", "a", "c", "d", "f", "k", "s"}


def clave_club(nombre: str | None) -> str:
    return " ".join(t for t in norm(nombre).split() if t not in _PALABRAS_VACIAS)


_CLAVES_CLUB = {clave_club(n): cat for cat, lista in CAT_CLUBES.items() for n in lista}
# variantes de nombre que usa Wikipedia
for _alias, _base in {"Internazionale": "Inter Milan", "Inter": "Inter Milan",
                      "Bayern München": "Bayern Munich", "PSG": "Paris Saint-Germain",
                      "Atlético de Madrid": "Atlético Madrid", "Milan": "AC Milan",
                      "Sporting Lisbon": "Sporting CP", "Club América": "América",
                      "Tigres": "Tigres UANL", "UNAM": "Pumas UNAM", "Al Hilal": "Al-Hilal",
                      "Al Nassr": "Al-Nassr", "Al Ittihad": "Al-Ittihad",
                      "Salzburg": "Red Bull Salzburg", "Shakhtar": "Shakhtar Donetsk",
                      "Crvena zvezda": "Red Star Belgrade", "FC København": "Copenhagen",
                      "SL Benfica": "Benfica", "AS Roma": "Roma", "SSC Napoli": "Napoli",
                      "Olympique de Marseille": "Marseille", "AS Monaco": "Monaco",
                      "Tottenham": "Tottenham Hotspur", "Newcastle": "Newcastle United",
                      "Bayer 04 Leverkusen": "Bayer Leverkusen", "BVB": "Borussia Dortmund",
                      "Borussia Dortmund II": None}.items():
    if _base:
        _CLAVES_CLUB.setdefault(clave_club(_alias), _CLAVES_CLUB[clave_club(_base)])


def categoria_club(nombre: str | None, wiki: str | None, liga: str | None) -> str:
    for n in (nombre, wiki):
        if n and (cat := _CLAVES_CLUB.get(clave_club(n))):
            return cat
    if liga in LIGAS_TOP5:
        return "3 Importante"
    return CAT_CLUB_RESTO


def categoria_seleccion(nombre: str) -> str:
    for cat, lista in CAT_SELECCIONES.items():
        if nombre in lista:
            return cat
    return CAT_SEL_RESTO


def actualizar_categorias(con) -> None:
    for s in con.execute("SELECT seleccion_id, nombre FROM selecciones").fetchall():
        con.execute("UPDATE selecciones SET categoria=? WHERE seleccion_id=?",
                    (categoria_seleccion(s["nombre"]), s["seleccion_id"]))
    for c in con.execute("SELECT club_id, nombre, wiki_titulo, liga FROM clubes").fetchall():
        con.execute("UPDATE clubes SET categoria=? WHERE club_id=?",
                    (categoria_club(c["nombre"], c["wiki_titulo"], c["liga"]), c["club_id"]))
    con.commit()


# ---------------------------------------------------------------------------
# Ligas de los clubes (infobox de Wikipedia)
# ---------------------------------------------------------------------------
def _liga_de_wikitexto(texto: str) -> str | None:
    """Lee el parámetro 'league' del infobox de un club."""
    m = re.search(r"\n\s*\|\s*league\s*=\s*(.+)", texto)
    if not m:
        return None
    valor = m.group(1)
    valor = re.split(r"<ref|<!--", valor)[0]
    enlace = re.search(r"\[\[([^\]|#]+)", valor)
    if enlace:
        liga = enlace.group(1).strip()
    else:
        liga = re.sub(r"\{\{[^}]*\}\}|<[^>]+>|'''?", "", valor).strip(" |")
    # '2025–26 Premier League' -> 'Premier League'
    liga = re.sub(r"^\d{4}(?:[–-]\d{2,4})?\s+", "", liga).strip()
    return liga or None


def cmd_ligas(args):
    """Busca la liga actual de cada club en el infobox de su artículo de Wikipedia."""
    with conectar() as con:
        filtro = "" if args.refrescar else "AND liga IS NULL"
        clubes = con.execute(f"SELECT club_id, nombre, wiki_titulo FROM clubes "
                             f"WHERE wiki_titulo IS NOT NULL {filtro}").fetchall()
        if not clubes:
            print("Todos los clubes ya tienen liga. Usa --refrescar para volver a consultarlas.")
        _wiki.headers["User-Agent"] = user_agent()
        lote, hechos, encontradas = 20, 0, 0
        for i in range(0, len(clubes), lote):
            grupo = clubes[i:i + lote]
            titulos = {c["wiki_titulo"] for c in grupo}
            clave = "ligas:" + "|".join(sorted(titulos))
            datos = cache_get(clave)
            if datos is None:
                r = _wiki.get(WIKI_API, timeout=60, params={
                    "action": "query", "prop": "revisions", "rvprop": "content",
                    "rvslots": "main", "titles": "|".join(titulos), "redirects": 1,
                    "format": "json", "formatversion": 2, "maxlag": 5})
                datos = r.json()
                time.sleep(WIKI_PAUSA_S)
                if "error" in datos:
                    print(f"  Wikipedia respondió: {datos['error'].get('info')}")
                    continue
                cache_put(clave, datos)
            q = datos.get("query", {})
            # seguir normalizaciones y redirecciones hasta el título final
            destino = {t: t for t in titulos}
            for paso in ("normalized", "redirects"):
                mapa = {x["from"]: x["to"] for x in q.get(paso, [])}
                destino = {t: mapa.get(d, d) for t, d in destino.items()}
            liga_de = {}
            for pg in q.get("pages", []):
                revs = pg.get("revisions") or []
                if revs:
                    texto = revs[0].get("slots", {}).get("main", {}).get("content", "")
                    liga_de[pg["title"]] = _liga_de_wikitexto(texto)
            for c in grupo:
                liga = liga_de.get(destino.get(c["wiki_titulo"]))
                con.execute("UPDATE clubes SET liga=? WHERE club_id=?", (liga or "", c["club_id"]))
                encontradas += bool(liga)
            hechos += len(grupo)
            print(f"  {hechos}/{len(clubes)} clubes revisados")
        actualizar_categorias(con)
        sin = con.execute("SELECT nombre FROM clubes WHERE liga IS NULL OR liga=''").fetchall()
    print(f"Liga encontrada para {encontradas} de {len(clubes)} clubes consultados.")
    if sin:
        print(f"{len(sin)} clubes quedaron sin liga (sin artículo o sin infobox): "
              + ", ".join(r[0] for r in sin[:20]) + (" ..." if len(sin) > 20 else ""))


# ---------------------------------------------------------------------------
# Hojas de análisis del Excel
# ---------------------------------------------------------------------------
JORNADAS = ["Jornada 1", "Jornada 2", "Jornada 3", "Dieciseisavos", "Octavos", "Cuartos",
            "Semifinal", "Tercer lugar", "Final"]
HOJAS_GENERALES = ["Leyenda", "Criterios", "Recambio", "Top jugadores", "Categorías", "Ligas",
                   "Clubes", "Top clubes por jornada", "Jugadores"]


def minutos_por_jornada(con) -> tuple[list[dict], set[str]]:
    """Una fila por convocado del Mundial con su club, liga, categorías y minutos
    en cada jornada. Devuelve también las jornadas que tienen alineaciones."""
    jornada, grupos = {}, {}
    for p in con.execute("SELECT * FROM partidos WHERE competicion=? ORDER BY fecha", (WC_ETIQUETA,)):
        r = ronda_es(p["ronda"])
        for sid in (p["local_id"], p["visita_id"]):
            if r.startswith("Grupo") or not r:
                grupos.setdefault(sid, []).append(p["partido_id"])
                jornada[(p["partido_id"], sid)] = f"Jornada {len(grupos[sid])}"
            else:
                jornada[(p["partido_id"], sid)] = r
    minutos, partidos_j, con_datos = {}, {}, set()
    for r in con.execute("""SELECT pa.partido_id, pa.seleccion_id, pa.jugador_id, pa.minutos
                            FROM participaciones pa JOIN partidos p USING(partido_id)
                            WHERE p.competicion=?""", (WC_ETIQUETA,)):
        j = jornada.get((r["partido_id"], r["seleccion_id"]))
        if j:
            con_datos.add(j)
            clave = (r["jugador_id"], r["seleccion_id"])
            minutos.setdefault(clave, {}).setdefault(j, 0)
            minutos[clave][j] += r["minutos"] or 0
    filas = []
    for r in con.execute("""SELECT c.jugador_id, l.seleccion_id, j.nombre, j.fecha_nac, c.posicion,
                                   s.nombre AS seleccion, s.categoria AS cat_sel,
                                   cl.nombre AS club, cl.pais_club, cl.liga, cl.categoria AS cat_club
                            FROM convocados c JOIN listas l USING(lista_id)
                            JOIN jugadores j USING(jugador_id)
                            JOIN selecciones s ON s.seleccion_id = l.seleccion_id
                            LEFT JOIN clubes cl USING(club_id)
                            WHERE l.ventana='mundial-2026'"""):
        m = minutos.get((r["jugador_id"], r["seleccion_id"]), {})
        edad = _edad(r["fecha_nac"], date(2026, 6, 11))
        filas.append({
            "jugador": r["nombre"], "seleccion": nombre_es(r["seleccion"]),
            "cat_sel": r["cat_sel"] or categoria_seleccion(r["seleccion"]),
            "club": r["club"] or "(sin club)",
            "cat_club": r["cat_club"] or CAT_CLUB_RESTO,
            "liga": r["liga"] or "(sin dato)",
            "pais_club": nombre_es(r["pais_club"]) if r["pais_club"] else "",
            "pos": r["posicion"], "edad": int(edad) if edad else None,
            "min": [m.get(j, 0) for j in JORNADAS],
        })
    return filas, con_datos


def _cuenta_txt(valores) -> str:
    cuenta = {}
    for v in valores:
        cuenta[v] = cuenta.get(v, 0) + 1
    return ", ".join(f"{k} ({n})" for k, n in sorted(cuenta.items(), key=lambda x: (-x[1], norm(x[0]))))


def hojas_analisis(wb, con, e: dict, top_n: int = 15):
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, PatternFill
    from openpyxl.utils import get_column_letter as L

    actualizar_categorias(con)
    filas, con_datos = minutos_por_jornada(con)
    if not filas:
        return
    sin_datos_txt = "Wikipedia no publica las alineaciones de esta ronda; los minutos aparecen en 0."
    izq = Alignment(vertical="top", wrap_text=True)
    arriba_centro = Alignment(horizontal="center", vertical="top")
    gris_claro = PatternFill("solid", start_color="F2F2F2", end_color="F2F2F2")
    COLORES_CAT = {"1": "FCE4D6", "2": "FFF2CC", "3": "E2EFDA", "4": "FFFFFF"}
    fill_cat = lambda cat: PatternFill("solid", start_color=COLORES_CAT.get((cat or "4")[0], "FFFFFF"),
                                       end_color=COLORES_CAT.get((cat or "4")[0], "FFFFFF"))
    NUM = "#,##0;-#,##0;-"

    def cabecera(ws, fila, textos, col0=1):
        for k, t in enumerate(textos):
            c = ws.cell(fila, col0 + k, t)
            c.font, c.fill, c.alignment = e["f_cabecera"], e["R_CAB"], e["centro"]
            if t in JORNADAS and t not in con_datos:
                c.comment = Comment(sin_datos_txt, "selecciones.py")

    def celda(ws, f, c, v, num=False, negrita=False, alin=None):
        x = ws.cell(f, c, v)
        x.font = e["f_negrita"] if negrita else e["f_normal"]
        x.border = e["borde"]
        if num:
            x.alignment, x.number_format = e["centro"], NUM
        elif alin:
            x.alignment = alin
        return x

    def nueva_hoja(nombre):
        ws = wb.create_sheet(nombre)
        orden = [h for h in HOJAS_GENERALES if h in wb.sheetnames]
        # reacomoda: hojas generales primero, en el orden de HOJAS_GENERALES
        wb._sheets.sort(key=lambda h: (orden.index(h.title) if h.title in orden else len(orden)))
        return ws

    # ===== Jugadores (hoja base de todas las fórmulas) =====
    filas.sort(key=lambda f: (-sum(f["min"]), norm(f["jugador"])))
    jug = nueva_hoja("Jugadores")
    cab = ["Jugador", "Selección", "Categoría selección", "Club", "Categoría club", "Liga",
           "País del club", "Pos.", "Edad"] + JORNADAS + ["Total minutos"]
    cabecera(jug, 1, cab)
    J0 = 10                                   # primera jornada (columna J)
    JN = J0 + len(JORNADAS) - 1               # última jornada (columna R)
    CT = JN + 1                               # total (columna S)
    n = len(filas) + 1
    for i, f in enumerate(filas, start=2):
        celda(jug, i, 1, f["jugador"], negrita=True)
        celda(jug, i, 2, f["seleccion"])
        celda(jug, i, 3, f["cat_sel"]).fill = fill_cat(f["cat_sel"])
        celda(jug, i, 4, f["club"])
        celda(jug, i, 5, f["cat_club"]).fill = fill_cat(f["cat_club"])
        celda(jug, i, 6, f["liga"])
        celda(jug, i, 7, f["pais_club"])
        celda(jug, i, 8, f["pos"], alin=e["centro"])
        celda(jug, i, 9, f["edad"], num=True)
        for k, v in enumerate(f["min"]):
            celda(jug, i, J0 + k, v, num=True)
        celda(jug, i, CT, f"=SUM({L(J0)}{i}:{L(JN)}{i})", num=True, negrita=True)
    for j, w in enumerate([24, 16, 15, 22, 17, 20, 14, 5, 5] + [10] * len(JORNADAS) + [9], start=1):
        jug.column_dimensions[L(j)].width = w
    jug.row_dimensions[1].height = 32
    jug.freeze_panes = "B2"
    jug.auto_filter.ref = f"A1:{L(CT)}{n}"

    R = lambda col: f"Jugadores!${L(col)}$2:${L(col)}${n}"
    R_JUG, R_SEL, R_CSEL, R_CLUB, R_CCLUB, R_LIGA, R_PAIS = (R(c) for c in range(1, 8))
    R_TOT = R(CT)
    total_min = sum(sum(f["min"]) for f in filas)

    # ===== Clubes =====
    clubes = {}
    for f in filas:
        clubes.setdefault((f["club"], f["pais_club"]), []).append(f)
    orden = sorted(clubes.items(), key=lambda kv: (-sum(sum(f["min"]) for f in kv[1]), norm(kv[0][0])))
    ws = nueva_hoja("Clubes")
    cab = ["Club", "Categoría", "Liga", "País del club", "Jugadores en el Mundial",
           "Selecciones de sus jugadores"] + JORNADAS + ["Total minutos"]
    cabecera(ws, 1, cab)
    C0 = 7
    for i, ((club, pais), js) in enumerate(orden, start=2):
        celda(ws, i, 1, club, negrita=True)
        celda(ws, i, 2, js[0]["cat_club"]).fill = fill_cat(js[0]["cat_club"])
        celda(ws, i, 3, js[0]["liga"])
        celda(ws, i, 4, pais)
        celda(ws, i, 5, f"=COUNTIFS({R_CLUB},$A{i},{R_PAIS},$D{i})", num=True)
        celda(ws, i, 6, _cuenta_txt(f["seleccion"] for f in js), alin=izq)
        for k in range(len(JORNADAS)):
            celda(ws, i, C0 + k, f"=SUMIFS({R(J0 + k)},{R_CLUB},$A{i},{R_PAIS},$D{i})", num=True)
        celda(ws, i, C0 + len(JORNADAS), f"=SUM({L(C0)}{i}:{L(C0 + len(JORNADAS) - 1)}{i})",
              num=True, negrita=True)
    for j, w in enumerate([24, 17, 20, 14, 10, 32] + [10] * len(JORNADAS) + [10], start=1):
        ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{L(len(cab))}{len(orden) + 1}"

    # ===== Ligas =====
    ligas = {}
    for f in filas:
        ligas.setdefault(f["liga"], []).append(f)
    orden_l = sorted(ligas.items(), key=lambda kv: (-sum(sum(f["min"]) for f in kv[1]), norm(kv[0])))
    ws = nueva_hoja("Ligas")
    ws.cell(1, 1, "Ligas donde juegan los convocados del Mundial 2026").font = e["f_titulo"]
    ws.cell(2, 1, "Ordenadas por minutos totales. Usa el filtro para ordenar por jugadores o por "
                  "cualquier ronda. La liga sale del infobox del club en Wikipedia.").font = e["f_gris"]
    cats_sel = list(CAT_SELECCIONES) + [CAT_SEL_RESTO]
    cab = (["Liga", "País (de sus clubes)", "Clubes", "Jugadores"] +
           [f"Jugadores de selecc. {c[2:]}" for c in cats_sel] +
           ["Selecciones que aportan más jugadores"] +
           JORNADAS[:3] + ["Fase de grupos"] + JORNADAS[3:] + ["Eliminatorias", "Total minutos",
                                                               "% de los minutos del Mundial"])
    F0 = 4
    cabecera(ws, F0, cab)
    c_sel = 5 + len(cats_sel)                 # columna de texto de selecciones
    c_j = c_sel + 1                           # Jornada 1
    for i, (liga, js) in enumerate(orden_l, start=F0 + 1):
        celda(ws, i, 1, liga, negrita=True)
        celda(ws, i, 2, _cuenta_txt(f["pais_club"] for f in js if f["pais_club"]).split(" (")[0])
        celda(ws, i, 3, len({f["club"] for f in js}), num=True)
        celda(ws, i, 4, f"=COUNTIF({R_LIGA},$A{i})", num=True)
        for k, cat in enumerate(cats_sel):
            celda(ws, i, 5 + k, f'=COUNTIFS({R_LIGA},$A{i},{R_CSEL},"{cat}")', num=True)
        celda(ws, i, c_sel, _cuenta_txt(f["seleccion"] for f in js), alin=izq)
        col = c_j
        for k in range(3):
            celda(ws, i, col, f"=SUMIFS({R(J0 + k)},{R_LIGA},$A{i})", num=True); col += 1
        celda(ws, i, col, f"=SUM({L(c_j)}{i}:{L(c_j + 2)}{i})", num=True, negrita=True); col += 1
        ko0 = col
        for k in range(3, len(JORNADAS)):
            celda(ws, i, col, f"=SUMIFS({R(J0 + k)},{R_LIGA},$A{i})", num=True); col += 1
        celda(ws, i, col, f"=SUM({L(ko0)}{i}:{L(col - 1)}{i})", num=True, negrita=True); col += 1
        celda(ws, i, col, f"={L(c_j + 3)}{i}+{L(col - 1)}{i}", num=True, negrita=True)
        c_tot = col; col += 1
        x = celda(ws, i, col, f"=IFERROR({L(c_tot)}{i}/SUM({R_TOT}),0)", num=True)
        x.number_format = "0.0%"
    ult_l = F0 + len(orden_l)
    for j, w in enumerate([24, 14, 7, 9] + [11] * len(cats_sel) + [34] + [10] * (len(JORNADAS) + 2)
                          + [10, 11], start=1):
        ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[F0].height = 45
    ws.freeze_panes = ws.cell(F0 + 1, 2)
    ws.auto_filter.ref = f"A{F0}:{L(len(cab))}{ult_l}"

    # ===== Categorías: cruce club × selección =====
    ws = nueva_hoja("Categorías")
    cats_club = list(CAT_CLUBES) + [CAT_CLUB_RESTO]
    ws.cell(1, 1, "Cruce entre la categoría del club y la de la selección").font = e["f_titulo"]
    ws.cell(2, 1, "Los criterios de cada categoría están en la hoja Criterios.").font = e["f_gris"]
    fila = 4
    for titulo, rango_valor, fmt in (("Jugadores convocados", None, NUM),
                                     ("Minutos jugados en el Mundial", R_TOT, NUM),
                                     ("Porcentaje de los minutos del Mundial", R_TOT, "0.0%")):
        ws.cell(fila, 1, titulo).font = e["f_titulo"]
        fila += 1
        cabecera(ws, fila, ["Club \\ Selección"] + [c[2:] for c in cats_sel] + ["Total"])
        f_ini = fila + 1
        for a, cc in enumerate(cats_club):
            r = f_ini + a
            celda(ws, r, 1, cc[2:], negrita=True).fill = fill_cat(cc)
            for b, cs in enumerate(cats_sel):
                if rango_valor is None:
                    v = f'=COUNTIFS({R_CCLUB},"{cc}",{R_CSEL},"{cs}")'
                else:
                    v = f'=SUMIFS({rango_valor},{R_CCLUB},"{cc}",{R_CSEL},"{cs}")'
                    if fmt.endswith("%"):
                        v = f"=IFERROR({v[1:]}/SUM({R_TOT}),0)"
                celda(ws, r, 2 + b, v, num=True).number_format = fmt
            celda(ws, r, 2 + len(cats_sel), f"=SUM({L(2)}{r}:{L(1 + len(cats_sel))}{r})",
                  num=True, negrita=True).number_format = fmt
        r = f_ini + len(cats_club)
        celda(ws, r, 1, "Total", negrita=True)
        for b in range(len(cats_sel) + 1):
            celda(ws, r, 2 + b, f"=SUM({L(2 + b)}{f_ini}:{L(2 + b)}{r - 1})",
                  num=True, negrita=True).number_format = fmt
        fila = r + 3
    ws.column_dimensions["A"].width = 22
    for b in range(len(cats_sel) + 1):
        ws.column_dimensions[L(2 + b)].width = 15

    # ===== Top jugadores =====
    tj = nueva_hoja("Top jugadores")
    tj.cell(1, 1, "Jugadores con más minutos en el Mundial 2026").font = e["f_titulo"]
    tj.cell(2, 1, "El orden se calculó al generar el archivo; los minutos son fórmulas sobre la "
                  "hoja Jugadores, que tiene la lista completa con filtros.").font = e["f_gris"]
    elite_sel, elite_club = cats_sel[0], cats_club[0]
    bloques = [
        ("Todos los jugadores", lambda f: True),
        (f"Juegan en clubes {elite_club[2:]}", lambda f: f["cat_club"] == elite_club),
        (f"Juegan en clubes {cats_club[1][2:]}", lambda f: f["cat_club"] == cats_club[1]),
        (f"Juegan en selecciones {elite_sel[2:]}", lambda f: f["cat_sel"] == elite_sel),
        (f"Clubes {elite_club[2:]} y selecciones {elite_sel[2:]}",
         lambda f: f["cat_club"] == elite_club and f["cat_sel"] == elite_sel),
        (f"Clubes {elite_club[2:]} en selecciones que NO son {elite_sel[2:]}",
         lambda f: f["cat_club"] == elite_club and f["cat_sel"] != elite_sel),
        (f"Fuera de las cinco grandes ligas", lambda f: f["liga"] not in LIGAS_TOP5),
    ]
    fila = 4
    for titulo, cond in bloques:
        tj.cell(fila, 1, titulo).font = e["f_titulo"]
        fila += 1
        cabecera(tj, fila, ["#", "Jugador", "Selección", "Categoría selección", "Club",
                            "Categoría club", "Liga", "Pos.", "Edad", "Minutos"])
        sel = [f for f in filas if cond(f) and sum(f["min"]) > 0][:top_n * 2]
        if not sel:
            tj.cell(fila + 1, 1, "Ningún jugador cumple esta condición.").font = e["f_gris"]
            fila += 3
            continue
        for pos, f in enumerate(sel, start=1):
            fila += 1
            vals = [pos, f["jugador"], f["seleccion"], f["cat_sel"], f["club"], f["cat_club"],
                    f["liga"], f["pos"], f["edad"]]
            for j, v in enumerate(vals, start=1):
                x = celda(tj, fila, j, v, negrita=(j == 2),
                          alin=arriba_centro if j in (1, 8, 9) else None)
                if j in (4, 6):
                    x.fill = fill_cat(v)
                elif pos % 2 == 0:
                    x.fill = gris_claro
            celda(tj, fila, 10, f"=SUMIFS({R_TOT},{R_JUG},B{fila},{R_SEL},C{fila})",
                  num=True, negrita=True)
        fila += 3
    for j, w in enumerate([5, 24, 16, 15, 22, 17, 20, 5, 5, 9], start=1):
        tj.column_dimensions[L(j)].width = w

    # ===== Top clubes por jornada =====
    tp = nueva_hoja("Top clubes por jornada")
    tp.cell(1, 1, f"Los {top_n} clubes cuyos jugadores sumaron más minutos en cada jornada "
                  "del Mundial 2026").font = e["f_titulo"]
    tp.cell(2, 1, "El orden se calculó al generar el archivo; los minutos son fórmulas "
                  "sobre la hoja Jugadores.").font = e["f_gris"]
    fila = 4
    for k, jn in enumerate(JORNADAS):
        tp.cell(fila, 1, jn).font = e["f_titulo"]
        if jn not in con_datos:
            tp.cell(fila + 1, 1, sin_datos_txt).font = e["f_gris"]
            fila += 3
            continue
        fila += 1
        cabecera(tp, fila, ["#", "Club", "Categoría", "Liga", "País del club", "Minutos",
                            "Jugadores que sumaron minutos", "Selecciones de esos jugadores"])
        ranking = []
        for (club, pais), js in clubes.items():
            jugaron = [f for f in js if f["min"][k] > 0]
            if jugaron:
                ranking.append((sum(f["min"][k] for f in jugaron), club, pais, jugaron))
        ranking.sort(key=lambda x: (-x[0], norm(x[1])))
        for pos, (_, club, pais, jugaron) in enumerate(ranking[:top_n], start=1):
            fila += 1
            jugaron.sort(key=lambda f: -f["min"][k])
            vals = [pos, club, jugaron[0]["cat_club"], jugaron[0]["liga"], pais,
                    f"=SUMIFS({R(J0 + k)},{R_CLUB},B{fila},{R_PAIS},E{fila})",
                    ", ".join(f"{f['jugador']} ({f['seleccion']}, {f['min'][k]}')" for f in jugaron),
                    _cuenta_txt(f["seleccion"] for f in jugaron)]
            for j, v in enumerate(vals, start=1):
                x = celda(tp, fila, j, v, num=(j == 6), negrita=(j in (2, 6)),
                          alin=arriba_centro if j == 1 else izq)
                if j == 3:
                    x.fill = fill_cat(v)
                elif pos % 2 == 0:
                    x.fill = gris_claro
        fila += 3
    for j, w in enumerate([5, 24, 17, 20, 14, 9, 70, 30], start=1):
        tp.column_dimensions[L(j)].width = w

    # ===== Criterios =====
    cr = nueva_hoja("Criterios")
    cr.cell(1, 1, "Criterios de categorías").font = e["f_titulo"]
    textos = [
        "Las categorías son una clasificación propia para agrupar, no un ranking oficial. "
        "Se editan al inicio de selecciones.py (CAT_SELECCIONES y CAT_CLUBES).",
        "Selecciones. 1 Élite: campeones del mundo recientes o semifinalistas habituales. "
        "2 Segunda línea: llegan seguido a eliminatorias del Mundial o ganan su confederación. "
        "3 Buen equipo: clasifican con regularidad a Mundiales o finales continentales. 4 Resto.",
        "Clubes. 1 Top 10 mundial: combinan títulos de Champions recientes, coeficiente UEFA, "
        "presupuesto e historia. 2 Top segunda línea: candidatos habituales a cuartos de Champions "
        "y gigantes de Sudamérica. 3 Importante: cualquier club de primera división en Inglaterra, "
        "España, Italia, Alemania o Francia, más los clubes dominantes de otras ligas fuertes. 4 Resto.",
    ]
    for k, t in enumerate(textos, start=2):
        x = cr.cell(k, 1, t)
        x.font, x.alignment = e["f_normal"], izq
        cr.merge_cells(start_row=k, start_column=1, end_row=k, end_column=6)
        cr.row_dimensions[k].height = 60
    fila = 6
    cabecera(cr, fila, ["Selección", "Categoría"])
    sels = con.execute("SELECT nombre, categoria FROM selecciones").fetchall()
    for i, s in enumerate(sorted(sels, key=lambda s: (s["categoria"], norm(nombre_es(s["nombre"])))),
                          start=fila + 1):
        celda(cr, i, 1, nombre_es(s["nombre"]))
        celda(cr, i, 2, s["categoria"]).fill = fill_cat(s["categoria"])
    cabecera(cr, fila, ["Club", "Categoría", "Liga"], col0=4)
    cls = con.execute("""SELECT DISTINCT cl.nombre, cl.categoria, cl.liga FROM clubes cl
                         WHERE cl.categoria <> ? ORDER BY cl.categoria, cl.nombre""",
                      (CAT_CLUB_RESTO,)).fetchall()
    for i, c in enumerate(cls, start=fila + 1):
        celda(cr, i, 4, c["nombre"])
        celda(cr, i, 5, c["categoria"]).fill = fill_cat(c["categoria"])
        celda(cr, i, 6, c["liga"])
    x = cr.cell(fila + len(cls) + 2, 4, "Los clubes que no aparecen aquí están en 4 Resto.")
    x.font = e["f_gris"]
    for j, w in enumerate([20, 17, 3, 26, 20, 22], start=1):
        cr.column_dimensions[L(j)].width = w


def cmd_reporte(args):
    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("Falta una librería. Instálala con:  pip install openpyxl")
        return

    FUENTE = "Arial"
    f_normal, f_negrita = Font(name=FUENTE, size=10), Font(name=FUENTE, size=10, bold=True)
    f_titulo = Font(name=FUENTE, size=13, bold=True)
    f_gris = Font(name=FUENTE, size=10, color="808080")
    f_cabecera = Font(name=FUENTE, size=9, bold=True, color="FFFFFF")
    relleno = lambda c: PatternFill("solid", start_color=c, end_color=c)
    R_CAB, R_CAB2 = relleno("1F3864"), relleno("2F5496")
    R_TIT, R_SUP, R_BAN = relleno("C6EFCE"), relleno("DDEBF7"), relleno("FFF2CC")
    R_NC, R_SD = relleno("EDEDED"), relleno("FFFFFF")
    borde = Border(bottom=Side(style="thin", color="D9D9D9"))
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # minutos>0 titular, minutos<0 entró de cambio, 0 convocado sin jugar
    FORMATO_MIN = '"T "0"\'";"S "0"\'";"Banca"'

    nombre_archivo = f"recambio_{args.ventana}.xlsx" if args.ventana else "recambio_mundial.xlsx"
    salida = BASE_DIR / nombre_archivo
    wb = Workbook()

    with conectar() as con:
        actualizar_categorias(con)
        # ---------------- Hoja Leyenda ----------------
        ws = wb.active
        ws.title = "Leyenda"
        filas_ley = [
            ("Cómo leer las hojas por selección", None, f_titulo),
            ("Cada fila es un jugador (orden alfabético) y cada columna un partido o una convocatoria.", None, f_normal),
            ("", None, f_normal),
            ("T 90'", R_TIT, f_normal, "Titular; minutos jugados"),
            ("S 24'", R_SUP, f_normal, "Entró de cambio; minutos jugados"),
            ("Banca", R_BAN, f_normal, "Convocado, no jugó"),
            ("No conv.", R_NC, f_gris, "No fue convocado"),
            ("Sin datos", R_SD, f_gris, "Estaba en la lista, pero Wikipedia no publica la alineación de ese partido"),
            ("Convocado", R_BAN, f_normal, "Columna de convocatoria sin partidos cargados todavía"),
            ("", None, f_normal),
            ("Las columnas PJ, Titular, Cambio, Banca y Minutos son fórmulas sobre las celdas de partido.", None, f_normal),
            ("Los minutos se cuentan sobre 90 (120 con prórroga), sin tiempo añadido.", None, f_normal),
            ("Los encabezados tienen filtro: puedes ordenar por Minutos o por cualquier partido.", None, f_normal),
            ("", None, f_normal),
            ("Hojas de análisis del Mundial", None, f_titulo),
            ("Criterios: cómo se definió cada categoría de selección y de club, y en cuál quedó cada una.",
             None, f_normal),
            ("Top jugadores: los que más minutos jugaron, en general, en clubes top, en selecciones "
             "top y combinaciones.", None, f_normal),
            ("Categorías: cruce entre categoría del club y de la selección (jugadores y minutos).",
             None, f_normal),
            ("Ligas: cuántos jugadores aporta cada liga y cuántos minutos suman por ronda.", None, f_normal),
            ("Clubes: minutos de los jugadores de cada club en cada jornada, con sus selecciones. "
             "Club = el de la lista oficial de 26.", None, f_normal),
            ("Top clubes por jornada: los 15 clubes con más minutos en cada jornada, con nombre, "
             "selección y minutos de cada jugador.", None, f_normal),
            ("Jugadores: la lista completa con club, liga, categorías y minutos por jornada. "
             "De ella salen las fórmulas de las demás hojas.", None, f_normal),
            ("Jornada 1, 2 y 3 son los tres partidos de grupo de cada selección, en orden de fecha.",
             None, f_normal),
            ("Fuente: Wikipedia (CC BY-SA 4.0), páginas del Mundial 2026 y de cada selección. "
             f"Generado el {date.today().strftime('%d/%m/%Y')}.", None, f_gris),
        ]
        for i, fila in enumerate(filas_ley, start=1):
            c = ws.cell(i, 1, fila[0])
            c.font = fila[2]
            if fila[1] is not None:
                c.fill, c.alignment = fila[1], centro
                ws.cell(i, 2, fila[3]).font = f_normal
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 80

        # ---------------- Hoja Recambio ----------------
        if args.ventana:
            filas = datos_recambio(con, args.ventana)
            if not filas:
                print(f"No hay listas de '{args.ventana}'. Corre antes:  "
                      f"python selecciones.py snapshot --ventana {args.ventana}")
            else:
                wr = wb.create_sheet("Recambio", 1)
                cab = ["Selección", "Categoría", "Lista Mundial", f"Lista {args.ventana}", "Siguen", "Salen",
                       "Entran", "Retención", "Similitud (Jaccard)", "Edad media Mundial",
                       "Edad ponderada por minutos (Mundial)", f"Edad media {args.ventana}",
                       "Minutos en el Mundial de los que salen", "Salen", "Entran"]
                for j, t in enumerate(cab, start=1):
                    c = wr.cell(1, j, t)
                    c.font, c.fill, c.alignment = f_cabecera, R_CAB, centro
                for i, f in enumerate(filas, start=2):
                    vals = [f["seleccion"], f["categoria"], f["mundial"], f["ventana"], f["siguen"],
                            f["salen"], f["entran"], f"=IFERROR(E{i}/C{i},0)",
                            f"=IFERROR(E{i}/(C{i}+G{i}),0)",
                            f["edad_a"], f["edad_pond"], f["edad_b"], f["min_salen"],
                            f["salen_n"], f["entran_n"]]
                    for j, v in enumerate(vals, start=1):
                        c = wr.cell(i, j, v)
                        c.font, c.border = f_normal, borde
                        if j in (14, 15):
                            c.alignment = Alignment(wrap_text=True, vertical="top")
                    wr.cell(i, 8).number_format = "0.0%"
                    wr.cell(i, 9).number_format = "0.00"
                    wr.cell(i, 13).number_format = "#,##0"
                anchos = [22, 15, 9, 9, 8, 8, 8, 10, 11, 10, 13, 10, 13, 50, 50]
                for j, w in enumerate(anchos, start=1):
                    wr.column_dimensions[get_column_letter(j)].width = w
                wr.row_dimensions[1].height = 45
                wr.freeze_panes = "B2"
                wr.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{len(filas) + 1}"
                wr.cell(1, 9).comment = Comment(
                    "Jugadores que siguen / total de jugadores distintos entre ambas listas. "
                    "1 = misma lista; 0 = ninguno repite.", "selecciones.py")

        # ---------------- Una hoja por selección ----------------
        selecciones = con.execute("""SELECT * FROM selecciones WHERE seleccion_id IN
                                     (SELECT seleccion_id FROM listas)""").fetchall()
        selecciones = sorted(selecciones, key=lambda s: norm(nombre_es(s["nombre"])))
        usados = set()
        for s in selecciones:
            sid = s["seleccion_id"]
            listas = {r["ventana"]: r["lista_id"] for r in con.execute(
                "SELECT ventana, lista_id FROM listas WHERE seleccion_id=?", (sid,))}
            en_lista = {}
            for r in con.execute("""SELECT l.ventana, c.jugador_id FROM convocados c
                                    JOIN listas l USING(lista_id) WHERE l.seleccion_id=?""", (sid,)):
                en_lista.setdefault(r["ventana"], set()).add(r["jugador_id"])
            partidos = con.execute(
                """SELECT p.*, sl.nombre AS local, sv.nombre AS visita FROM partidos p
                   JOIN selecciones sl ON sl.seleccion_id=p.local_id
                   JOIN selecciones sv ON sv.seleccion_id=p.visita_id
                   WHERE ? IN (p.local_id, p.visita_id) ORDER BY p.fecha""", (sid,)).fetchall()
            part = {}
            for r in con.execute("""SELECT pa.* FROM participaciones pa WHERE pa.seleccion_id=?""", (sid,)):
                part[(r["partido_id"], r["jugador_id"])] = r
            con_alin = {pid for pid, _ in part}

            # columnas: partidos en orden de fecha; luego listas sin partidos
            columnas = []
            ventanas_con_partido = set()
            for p in partidos:
                v = _ventana_de_partido(p["competicion"], p["fecha"])
                ventanas_con_partido.add(v)
                rival = p["visita"] if p["local_id"] == sid else p["local"]
                comp = "FWC 2026" if p["competicion"] == WC_ETIQUETA else (p["competicion"] or "")
                fecha_txt = (datetime.strptime(p["fecha"], "%Y-%m-%d").strftime("%d/%m/%Y")
                             if p["fecha"] else "")
                titulo = f"{comp} {ronda_es(p['ronda'])}\n{fecha_txt}\nvs {nombre_es(rival)}"
                if p["marcador"] and (m := re.search(r"(\d+)\s*[–-]\s*(\d+)", p["marcador"])):
                    gl, gv = m.groups()
                    marc = f"{gl}-{gv}" if p["local_id"] == sid else f"{gv}-{gl}"
                    if p["prorroga"]:
                        marc += " t.e."
                    titulo += f" ({marc})"
                columnas.append(("partido", p, v, titulo))
            for v in sorted(listas):
                if v not in ventanas_con_partido and v != "mundial-2026":
                    columnas.append(("lista", None, v, f"Convocatoria\n{v}"))

            # filas: todos los jugadores que aparecen en listas o partidos
            jug_ids = set().union(*en_lista.values()) | {j for _, j in part}
            jugadores = con.execute(
                f"SELECT * FROM jugadores WHERE jugador_id IN ({','.join('?' * len(jug_ids))})",
                tuple(jug_ids)).fetchall()
            jugadores = sorted(jugadores, key=lambda j: norm(j["nombre"]))
            # posición y club de la lista más reciente donde aparece
            info = {}
            for r in con.execute("""SELECT c.jugador_id, c.posicion, cl.nombre AS club, cl.liga,
                                           cl.categoria AS cat_club, l.obtenida_en
                                    FROM convocados c JOIN listas l USING(lista_id)
                                    LEFT JOIN clubes cl USING(club_id)
                                    WHERE l.seleccion_id=? ORDER BY l.obtenida_en""", (sid,)):
                info[r["jugador_id"]] = r

            hoja = nombre_es(s["nombre"])
            hoja = re.sub(r"[\[\]:*?/\\]", "", hoja)[:31]
            while hoja in usados:
                hoja = hoja[:29] + "_2"
            usados.add(hoja)
            ws = wb.create_sheet(hoja)

            fijas = ["Jugador", "Pos.", "Nacimiento", "Club", "Liga", "Categoría del club",
                     "PJ", "Titular", "Cambio", "Banca", "Minutos"]
            C0 = len(fijas) + 1                      # primera columna de partidos
            ult = C0 + len(columnas) - 1
            for j, t in enumerate(fijas, start=1):
                c = ws.cell(1, j, t)
                c.font, c.fill, c.alignment = f_cabecera, R_CAB, centro
            for k, (tipo, p, v, titulo) in enumerate(columnas):
                c = ws.cell(1, C0 + k, titulo)
                c.font, c.fill, c.alignment = f_cabecera, (R_CAB2 if tipo == "partido" else R_CAB), centro
                if tipo == "partido" and p["partido_id"] not in con_alin:
                    c.comment = Comment("Wikipedia no publica la alineación de este partido.",
                                        "selecciones.py")

            for i, jg in enumerate(jugadores, start=2):
                jid = jg["jugador_id"]
                inf = info.get(jid)
                ws.cell(i, 1, jg["nombre"]).font = f_negrita
                ws.cell(i, 2, inf["posicion"] if inf else None).font = f_normal
                c = ws.cell(i, 3, date.fromisoformat(jg["fecha_nac"]) if jg["fecha_nac"] else None)
                c.font, c.number_format = f_normal, "dd/mm/yyyy"
                ws.cell(i, 4, inf["club"] if inf else None).font = f_normal
                ws.cell(i, 5, inf["liga"] if inf else None).font = f_normal
                ws.cell(i, 6, inf["cat_club"] if inf else None).font = f_normal
                for k, (tipo, p, v, _) in enumerate(columnas):
                    c = ws.cell(i, C0 + k)
                    c.alignment, c.font = centro, f_normal
                    convocado = jid in en_lista.get(v, set())
                    if tipo == "lista":
                        c.value, c.fill = ("Convocado", R_BAN) if convocado else ("No conv.", R_NC)
                        if not convocado:
                            c.font = f_gris
                        continue
                    pid = p["partido_id"]
                    r = part.get((pid, jid))
                    if r is not None:
                        mins = r["minutos"] or 0
                        if r["titular"]:
                            c.value, c.fill = max(mins, 1), R_TIT
                        elif mins > 0:
                            c.value, c.fill = -mins, R_SUP
                        else:
                            c.value, c.fill = 0, R_BAN
                        c.number_format = FORMATO_MIN
                        notas = []
                        if r["amarillas"]:
                            notas.append("Amarilla" if r["amarillas"] == 1 else f"{r['amarillas']} amarillas")
                        if r["expulsado"]:
                            notas.append("Expulsado")
                        if notas:
                            c.comment = Comment(", ".join(notas), "selecciones.py")
                    elif pid in con_alin:
                        if convocado:
                            c.value, c.fill, c.number_format = 0, R_BAN, FORMATO_MIN
                        elif v in listas:
                            c.value, c.fill, c.font = "No conv.", R_NC, f_gris
                        else:   # no hay lista guardada para esa fecha: no sabemos si fue llamado
                            c.value, c.fill, c.font = "No jugó", R_NC, f_gris
                    else:
                        c.value, c.fill, c.font = (("Sin datos", R_SD, f_gris) if convocado
                                                   else ("No conv.", R_NC, f_gris))
                # resumen con fórmulas sobre las celdas de partido
                if columnas:
                    rng = f"{get_column_letter(C0)}{i}:{get_column_letter(ult)}{i}"
                    ws.cell(i, 8, f'=COUNTIF({rng},">0")')
                    ws.cell(i, 9, f'=COUNTIF({rng},"<0")')
                    ws.cell(i, 7, f"=H{i}+I{i}")
                    ws.cell(i, 10, f"=COUNTIF({rng},0)")
                    ws.cell(i, 11, f'=SUMIF({rng},">0")-SUMIF({rng},"<0")')
                for j in range(1, C0):
                    c = ws.cell(i, j)
                    c.border = borde
                    if j >= 7:
                        c.font, c.alignment = f_normal, centro
                        c.number_format = "#,##0"

            ws.row_dimensions[1].height = 62
            for j, w in enumerate([26, 6, 11, 22, 18, 17, 5, 7, 7, 7, 8], start=1):
                ws.column_dimensions[get_column_letter(j)].width = w
            for k in range(len(columnas)):
                ws.column_dimensions[get_column_letter(C0 + k)].width = 17
            ws.freeze_panes = ws.cell(2, C0)
            ws.auto_filter.ref = f"A1:{get_column_letter(max(ult, C0 - 1))}{len(jugadores) + 1}"

    # ---------------- Análisis: jugadores, clubes, ligas y categorías ----------------
    hojas_analisis(wb, con, dict(f_normal=f_normal, f_negrita=f_negrita, f_titulo=f_titulo,
                              f_gris=f_gris, f_cabecera=f_cabecera, R_CAB=R_CAB, R_CAB2=R_CAB2,
                              borde=borde, centro=centro))

    wb.save(salida)
    hojas = sum(1 for h in wb.sheetnames if h not in HOJAS_GENERALES)
    print(f"Excel generado: {salida}")
    print("  Hojas generales: " + ", ".join(h for h in wb.sheetnames if h in HOJAS_GENERALES))
    print(f"  {hojas} hojas de selección")
    sin_liga = con.execute("SELECT COUNT(*) FROM clubes WHERE liga IS NULL").fetchone()[0]
    if sin_liga:
        print(f"  Aviso: {sin_liga} clubes sin liga. Corre 'python selecciones.py ligas' y vuelve "
              "a generar el reporte.")


# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("init", help="crea la base y pide tu correo").set_defaults(f=cmd_init)

    p = sub.add_parser("mundial", help="listas de 26 del Mundial 2026")
    p.add_argument("--refrescar", action="store_true", help="ignora la caché del día")
    p.set_defaults(f=cmd_mundial)

    p = sub.add_parser("partidos", help="alineaciones y minutos (por defecto, el Mundial 2026)")
    p.add_argument("--paginas", nargs="*", help="otras páginas de Wikipedia con partidos")
    p.add_argument("--competicion", help="etiqueta para esas páginas, p.ej. 'Nations League 2026-27'")
    p.add_argument("--muestra", action="store_true", help="imprime una alineación para revisarla")
    p.add_argument("--refrescar", action="store_true")
    p.set_defaults(f=cmd_partidos)

    p = sub.add_parser("ligas", help="busca la liga de cada club en Wikipedia")
    p.add_argument("--refrescar", action="store_true", help="vuelve a consultar todas")
    p.set_defaults(f=cmd_ligas)

    p = sub.add_parser("tablas", help="tabla de posiciones de cada liga (clasificación mundial de clubes)")
    p.add_argument("--ligas", nargs="*", help="solo estas ligas, p.ej. \"Liga MX\" \"Premier League\"")
    p.add_argument("--refrescar", action="store_true")
    p.set_defaults(f=lambda a: __import__("clasificacion").cmd_tablas(a))

    p = sub.add_parser("web", help="genera la carpeta docs/ (app instalable para GitHub Pages)")
    p.set_defaults(f=lambda a: __import__("clasificacion").cmd_web(a))

    p = sub.add_parser("panel", help="genera panel.html con la clasificación mundial y sus filtros")
    p.set_defaults(f=lambda a: __import__("clasificacion").cmd_panel(a))

    p = sub.add_parser("snapshot", help="fotografía la 'Current squad' de cada selección")
    p.add_argument("--ventana", required=True, help="etiqueta, p.ej. 2026-10")
    p.add_argument("--equipos", nargs="*", help="solo estas selecciones (nombre en inglés)")
    p.add_argument("--refrescar", action="store_true")
    p.set_defaults(f=cmd_snapshot)

    p = sub.add_parser("reporte", help="genera el Excel: una hoja por selección y el recambio")
    p.add_argument("--ventana", help="convocatoria a comparar con el Mundial, p.ej. 2026-10")
    p.set_defaults(f=cmd_reporte)

    args = ap.parse_args()
    args.f(args)


if __name__ == "__main__":
    main()
